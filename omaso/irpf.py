"""Planilha anual 'Resumo IRPF Rural' (abas Resumo, Lançamentos, Mês a mês, A identificar), com totais por fórmula."""
import os, re
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from .util import ABR, is_idn, notas_abertas, ativos, sem_acento

ESCURO = PatternFill('solid', fgColor='FF1D1B18')
CLARO = 'FFF4EFE6'
NUM = '#,##0.00;\\(#,##0.00\\);\\-'


def nome_arquivo(ie):
    if ie.get('nome_arquivo'):
        return ie['nome_arquivo']
    n = re.sub(r'\s+e\s+outr[oa]s?$', '', ie['nome'].strip(), flags=re.I)
    w = sem_acento(n).split()
    return '_'.join([w[0], w[-1]] if len(w) > 1 else w)


def _doc(ie):
    d = re.sub(r'\D', '', ie.get('doc', ''))
    if len(d) == 11:
        return f'CPF {d[:3]}.{d[3:6]}.{d[6:9]}-{d[9:]}'
    if len(d) == 14:
        return f'CNPJ {d[:2]}.{d[2:5]}.{d[5:8]}/{d[8:12]}-{d[12:]}'
    return ie.get('doc', '')


def _head(ws, row, cols):
    for i, v in enumerate(cols, 1):
        c = ws.cell(row=row, column=i, value=v)
        c.font = Font(bold=True, color=CLARO)
        c.fill = ESCURO
        c.alignment = Alignment(vertical='center', wrap_text=True)


def _desc(t):
    return (t.get('label') or '').replace('_', ' · ').replace('-', ' ')


def gerar(ie, ano, ts, out):
    """Grava <out>/<pasta da IE>/Resumo_IRPF_Rural/Resumo_IRPF_Rural_<Nome>_<ano>.xlsx; devolve (caminho relativo, Workbook)."""
    ano = str(ano)
    ts = sorted([t for t in ativos(ts) if (t.get('date') or '')[:4] == ano], key=lambda t: (t['date'], t.get('cat', '')))
    lanc = [t for t in ts if not is_idn(t)]
    idn = [t for t in ts if is_idn(t)]
    wb = Workbook()
    rs = wb.active
    rs.title = 'Resumo'
    wl = wb.create_sheet('Lançamentos')
    wm = wb.create_sheet('Mês a mês')
    wi = wb.create_sheet('A identificar')

    # Lançamentos
    _head(wl, 1, ['Data', 'Mês', 'Propriedade', 'Tipo', 'Categoria (Livro Caixa)', 'Descrição / documento', 'Cabeças', 'Valor (R$)', 'Pendência aberta', 'Pasta no Drive (dentro de INSCRIÇÕES)'])
    for r, t in enumerate(lanc, 2):
        tipo = 'Receita' if t['kind'] == 'venda' else 'Despesa'
        vals = [t['date'], int(t['date'][5:7]), t.get('prop'), tipo, t.get('cat'), _desc(t), t.get('heads') or None,
                t.get('value'), ' | '.join(notas_abertas(t)) or None, t.get('pasta')]
        for c, v in enumerate(vals, 1):
            wl.cell(row=r, column=c, value=v)
        wl.cell(row=r, column=8).number_format = NUM
    n = max(2, len(lanc) + 1)
    for col, w in zip('ABCDEFGHIJ', [11, 6, 22, 10, 36, 48, 9, 16, 60, 70]):
        wl.column_dimensions[col].width = w
    wl.freeze_panes = 'A2'
    wl.auto_filter.ref = f'A1:J{n}'
    H = lambda col: f'Lançamentos!${col}$2:${col}${n}'

    # Resumo
    datas = [t['date'] for t in ts]
    per = f"{datas[0][8:]}/{datas[0][5:7]} a {datas[-1][8:]}/{datas[-1][5:7]}/{ano}" if datas else 'sem lançamentos'
    rs['A1'] = f"Resumo IRPF Rural {ano} — {ie['nome']}"
    rs['A1'].font = Font(bold=True, size=14)
    rs['A2'] = f"{_doc(ie)} · IE {ie['ie']} · Livro Caixa da Atividade Rural (resultado de caixa: receitas recebidas − despesas pagas)"
    rs['A3'] = f'Gerado pelo OMASO a partir dos lançamentos organizados. Pagamentos "a identificar" ficam fora dos totais (aba A identificar). Período com dados: {per}.'
    _head(rs, 5, ['Categoria', 'Tipo', 'Lançamentos', 'Total (R$)'])
    cats = []
    for c in ['Venda de Bovinos', 'Compra de Animais Bovinos'] + sorted({t['cat'] for t in lanc} - {'Venda de Bovinos', 'Compra de Animais Bovinos'}):
        if any(t['cat'] == c for t in lanc):
            cats.append(c)
    r = 6
    for c in cats:
        tipo = 'Receita' if any(t['cat'] == c and t['kind'] == 'venda' for t in lanc) else 'Despesa'
        rs.cell(row=r, column=1, value=c)
        rs.cell(row=r, column=2, value=tipo)
        rs.cell(row=r, column=3, value=f'=COUNTIFS({H("E")},A{r})')
        rs.cell(row=r, column=4, value=f'=SUMIFS({H("H")},{H("E")},A{r})').number_format = NUM
        r += 1
    last = max(6, r - 1)
    r += 1
    t0 = r
    for lab, f in [('Total de receitas', f'=SUMIFS(D6:D{last},B6:B{last},"Receita")'),
                   ('Total de despesas (inclui compra de gado)', f'=SUMIFS(D6:D{last},B6:B{last},"Despesa")'),
                   ('Resultado da atividade rural (receitas − despesas)', f'=D{t0}-D{t0 + 1}'),
                   ('Resultado presumido (20% da receita bruta) — para comparação', f'=D{t0}*0.2')]:
        rs.cell(row=r, column=1, value=lab).font = Font(bold=True)
        c = rs.cell(row=r, column=4, value=f)
        c.number_format = NUM
        c.font = Font(bold=True)
        r += 1
    r += 1
    rs.cell(row=r, column=1, value='O resultado presumido de 20% da receita bruta é a opção prevista na legislação do IR para a atividade rural; a escolha entre resultado real e presumido deve ser confirmada com o contador.')
    r += 2
    _head(rs, r, ['Propriedade', 'Receitas (R$)', 'Despesas (R$)', 'Resultado (R$)'])
    for p in sorted({t['prop'] for t in lanc}):
        r += 1
        rs.cell(row=r, column=1, value=p)
        rs.cell(row=r, column=2, value=f'=SUMIFS({H("H")},{H("C")},A{r},{H("D")},"Receita")').number_format = NUM
        rs.cell(row=r, column=3, value=f'=SUMIFS({H("H")},{H("C")},A{r},{H("D")},"Despesa")').number_format = NUM
        rs.cell(row=r, column=4, value=f'=B{r}-C{r}').number_format = NUM
    rs.column_dimensions['A'].width = 58
    rs.column_dimensions['B'].width = 16
    rs.column_dimensions['C'].width = 14
    rs.column_dimensions['D'].width = 18

    # Mês a mês
    _head(wm, 1, ['Mês', 'Receitas (R$)', 'Despesas (R$)', 'Resultado (R$)', 'Resultado acumulado (R$)'])
    for m in range(1, 13):
        r = m + 1
        wm.cell(row=r, column=1, value=ABR[m])
        wm.cell(row=r, column=6, value=m)
        wm.cell(row=r, column=2, value=f'=SUMIFS({H("H")},{H("B")},F{r},{H("D")},"Receita")')
        wm.cell(row=r, column=3, value=f'=SUMIFS({H("H")},{H("B")},F{r},{H("D")},"Despesa")')
        wm.cell(row=r, column=4, value=f'=B{r}-C{r}')
        wm.cell(row=r, column=5, value='=D2' if m == 1 else f'=E{r - 1}+D{r}')
        for c in range(2, 6):
            wm.cell(row=r, column=c).number_format = NUM
    wm.cell(row=14, column=1, value='Total').font = Font(bold=True)
    for c, L in [(2, 'B'), (3, 'C'), (4, 'D')]:
        x = wm.cell(row=14, column=c, value=f'=SUM({L}2:{L}13)')
        x.number_format = NUM
        x.font = Font(bold=True)
    wm.column_dimensions['A'].width = 10
    for L in 'BCD':
        wm.column_dimensions[L].width = 18
    wm.column_dimensions['E'].width = 22
    wm.column_dimensions['F'].hidden = True

    # A identificar
    _head(wi, 1, ['Data', 'Propriedade', 'Descrição', 'Valor (R$)', 'Observação'])
    for r, t in enumerate(idn, 2):
        for c, v in enumerate([t['date'], t.get('prop'), _desc(t), t.get('value'), ' | '.join(t.get('notes') or []) or None], 1):
            wi.cell(row=r, column=c, value=v)
        wi.cell(row=r, column=4).number_format = NUM
    for col, w in zip('ABCDE', [11, 22, 48, 16, 60]):
        wi.column_dimensions[col].width = w

    rel = f"{ie['pasta']}/Resumo_IRPF_Rural/Resumo_IRPF_Rural_{nome_arquivo(ie)}_{ano}.xlsx"
    os.makedirs(os.path.dirname(os.path.join(out, rel)), exist_ok=True)
    wb.save(os.path.join(out, rel))
    return rel, wb
